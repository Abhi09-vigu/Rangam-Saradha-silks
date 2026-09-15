from django.contrib import admin
from django.contrib.auth.admin import UserAdmin
from .models import CustomUser, Address, Wishlist

from rangam_saradha_silk.admin import custom_admin_site

from django.utils.html import format_html
from django.urls import path
from django.http import HttpResponse
from django.utils import timezone
from django.core.exceptions import PermissionDenied
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_users_excel_response():
    """
    Generates and returns an HttpResponse containing an .xlsx Excel file with ALL
    users from the database, styled with custom headers, borders, zebra striping,
    and automatic column widths. Sensitive credentials (passwords, OTPs, tokens)
    are strictly excluded.
    """
    users = CustomUser.objects.all().order_by('-date_joined')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Users"
    ws.views.sheetView[0].showGridLines = True

    # 1. Header columns as specified
    headers = [
        "Username",
        "Email Address",
        "Phone Number",
        "Auth Provider",
        "Staff Status",
        "Superuser Status",
        "Active Status",
        "Date Joined",
        "Last Login",
    ]

    # Header styling matching Rangam Saradha Silk's burgundy theme
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="AF0446", end_color="AF0446", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )

    ws.row_dimensions[1].height = 28

    for col_idx, header_title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Safe datetime formatting helper
    def format_dt(dt):
        if not dt:
            return "-"
        try:
            if timezone.is_aware(dt):
                dt = timezone.localtime(dt)
            return dt.strftime("%Y-%m-%d %H:%M:%S")
        except Exception:
            return str(dt)

    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    alt_fill = PatternFill(start_color="FAF7F5", end_color="FAF7F5", fill_type="solid")

    for row_idx, user in enumerate(users, start=2):
        ws.row_dimensions[row_idx].height = 20
        is_alt = (row_idx % 2 == 1)

        # Safely resolve auth provider label
        auth_provider_label = "-"
        if user.auth_provider:
            if hasattr(user, 'get_auth_provider_display'):
                auth_provider_label = user.get_auth_provider_display() or user.auth_provider
            else:
                auth_provider_label = user.auth_provider

        row_data = [
            (user.username or "-", left_align),
            (user.email or "-", left_align),
            (str(user.phone_number) if user.phone_number else "-", center_align),
            (auth_provider_label, center_align),
            ("Yes" if user.is_staff else "No", center_align),
            ("Yes" if user.is_superuser else "No", center_align),
            ("Yes" if user.is_active else "No", center_align),
            (format_dt(user.date_joined), center_align),
            (format_dt(user.last_login) if user.last_login else "Never", center_align),
        ]

        for col_idx, (val, align) in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.alignment = align
            cell.border = thin_border
            if is_alt:
                cell.fill = alt_fill

    # Pin header row on scroll
    ws.freeze_panes = "A2"

    # Automatically set readable column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            cell_str = str(cell.value) if cell.value is not None else ""
            if len(cell_str) > max_len:
                max_len = len(cell_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="users.xlsx"'
    wb.save(response)
    return response


class CustomUserAdmin(UserAdmin):
    model = CustomUser
    change_list_template = "admin/accounts/customuser/change_list.html"
    list_display = ['username', 'email', 'phone_number', 'auth_provider', 'avatar_preview', 'is_verified', 'is_staff', 'date_joined', 'last_login']
    list_filter = UserAdmin.list_filter + ('auth_provider', 'is_verified')
    fieldsets = UserAdmin.fieldsets + (
        ('Custom Attributes', {'fields': ('phone_number', 'auth_provider', 'google_id', 'firebase_uid', 'profile_picture_url', 'is_verified', 'otp_code', 'otp_expiry')}),
    )
    add_fieldsets = UserAdmin.add_fieldsets + (
        ('Custom Attributes', {'fields': ('phone_number', 'auth_provider', 'is_verified')}),
    )

    def avatar_preview(self, obj):
        url = obj.get_avatar_url()
        if url:
            return format_html('<img src="{}" style="width: 32px; height: 32px; border-radius: 50%; object-fit: cover;" />', url)
        return "-"
    avatar_preview.short_description = "Avatar"

    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path(
                'export-excel/',
                self.admin_site.admin_view(self.export_all_users_excel),
                name='accounts_customuser_export_excel',
            ),
        ]
        return custom_urls + urls

    def export_all_users_excel(self, request):
        if not (request.user.is_authenticated and request.user.is_staff):
            raise PermissionDenied("Only authenticated admin staff can export user data.")
        return generate_users_excel_response()

class AddressAdmin(admin.ModelAdmin):
    list_display = ['user', 'full_name', 'phone_number', 'city', 'state', 'pincode', 'address_type', 'is_default']
    list_filter = ['state', 'address_type', 'is_default']
    search_fields = ['full_name', 'city', 'pincode', 'user__username']

class WishlistAdmin(admin.ModelAdmin):
    list_display = ['user', 'product', 'created_at']
    search_fields = ['user__username', 'product__name']

custom_admin_site.register(CustomUser, CustomUserAdmin)
custom_admin_site.register(Address, AddressAdmin)
custom_admin_site.register(Wishlist, WishlistAdmin)
